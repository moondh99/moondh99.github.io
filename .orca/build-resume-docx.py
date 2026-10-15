"""Build resume.docx from .orca/resume-table.html (same tables as resume.pdf).
Usage: python3 .orca/build-resume-docx.py  (run from repo root)"""
from bs4 import BeautifulSoup, NavigableString
import docx
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT as VA
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

soup = BeautifulSoup(open('.orca/resume-table.html', encoding='utf-8'), 'html.parser')
d = docx.Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = sec.top_margin = Mm(14); sec.bottom_margin = Mm(12)
FONT = 'Apple SD Gothic Neo'
st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(9.3)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.2
CONTENT_MM = 182

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hex_); tcPr.append(s)

def borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '10' if side == 'top' else '4'); e.set(qn('w:color'), '1A1A1A' if side == 'top' else 'B8B8B8')
        b.append(e)
    tblPr.append(b)

def runs(p, node, bold=False, small=False):
    for ch in node.children:
        if isinstance(ch, NavigableString):
            txt = str(ch)
            if txt.strip() or txt == ' ':
                r = p.add_run(txt.replace('\n', ''))
                r.bold = bold or None
                if small: r.font.size = Pt(8.6); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif ch.name == 'br': p.add_run().add_break()
        elif ch.name == 'b': runs(p, ch, True, small)
        else: runs(p, ch, bold, small)

def fill(cell, td):
    cell.paragraphs[0].text = ''
    p = cell.paragraphs[0]; first = True
    img = td.find('img')
    if img:
        p.add_run().add_picture(img['src'].lstrip('/'), width=Mm(29)); return
    for ch in td.children:
        if getattr(ch, 'name', None) == 'ul':
            for li in ch.find_all('li', recursive=False):
                q = p if first and not p.text else cell.add_paragraph()
                q.add_run('• '); runs(q, li); q.paragraph_format.left_indent = Mm(3); q.paragraph_format.first_line_indent = Mm(-3)
                first = False
        elif getattr(ch, 'name', None) == 'span':
            q = cell.add_paragraph(); runs(q, ch, small=True)
        elif getattr(ch, 'name', None) == 'br':
            p.add_run().add_break()
        elif getattr(ch, 'name', None) == 'b':
            runs(p, ch, True)
        elif isinstance(ch, NavigableString):
            if str(ch).strip(): p.add_run(str(ch))
        else:
            runs(p, ch)
    if td.name == 'th':
        for q in cell.paragraphs:
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in q.runs: r.bold = True
        shade(cell, 'F0F0F0')
    if 'c' in (td.get('class') or []):
        for q in cell.paragraphs: q.alignment = WD_ALIGN_PARAGRAPH.CENTER

def build(tbl):
    widths = [c.get('style', '') for c in tbl.find_all('col')]
    fixed = [float(w.split('width:')[1].rstrip('m;')) if 'width' in w else None for w in widths]
    free = (CONTENT_MM - sum(w for w in fixed if w)) / max(1, fixed.count(None))
    mm = [w or free for w in fixed]
    rows = tbl.find_all('tr'); n = len(mm)
    t = d.add_table(rows=len(rows), cols=n); t.autofit = False; borders(t)
    top = 'top' in (tbl.get('class') or [])
    occupied = set()
    for ri, tr in enumerate(rows):
        ci = 0
        for td in tr.find_all(['td', 'th'], recursive=False):
            while (ri, ci) in occupied: ci += 1
            cs, rs = int(td.get('colspan', 1)), int(td.get('rowspan', 1))
            cell = t.cell(ri, ci)
            if cs > 1 or rs > 1:
                cell = cell.merge(t.cell(ri + rs - 1, ci + cs - 1))
                for a in range(ri, ri + rs):
                    for b in range(ci, ci + cs): occupied.add((a, b))
            fill(cell, td)
            cell.vertical_alignment = VA.TOP if top and td.name == 'td' else VA.CENTER
            ci += cs
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    for i, col in enumerate(t.columns): col.width = Mm(mm[i])
    for row in t.rows:
        for i, c in enumerate(row.cells): c.width = Mm(mm[i])
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    if tbl.find('thead'):
        trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); trPr.append(h)

for el in soup.body.find_all(['h1', 'h2', 'table'], recursive=False):
    if el.name == 'h1':
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(el.get_text()); r.bold = True; r.font.size = Pt(20)
        p.paragraph_format.space_after = Pt(8)
    elif el.name == 'h2':
        p = d.add_paragraph(); r = p.add_run('▌' + el.get_text()); r.bold = True; r.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
    else:
        build(el)
d.core_properties.title = '문대한 이력서'; d.core_properties.author = '문대한'
d.save('resume.docx')
print('ok')
